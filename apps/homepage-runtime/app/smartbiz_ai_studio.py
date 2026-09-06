# SMARTBIZ_AI_STUDIO_PLATFORM_V075

from __future__ import annotations

import csv
import io
import json
import mimetypes
import zipfile
from pathlib import Path
from typing import Any, Callable
from uuid import uuid4

import httpx
from docx import Document
from fastapi import File, Request, UploadFile
from fastapi.responses import HTMLResponse, JSONResponse, RedirectResponse
from openpyxl import load_workbook
from pypdf import PdfReader


MAX_UPLOAD_BYTES = 25 * 1024 * 1024
MAX_EXTRACTED_CHARS = 16000

DOCUMENT_ASSISTANT_URL = (
    "https://botconnector.id/"
    "document-assistant/api/convert/handwriting"
)


def _items(payload: Any) -> list[dict[str, Any]]:
    if isinstance(payload, list):
        return [
            item
            for item in payload
            if isinstance(item, dict)
        ]

    if isinstance(payload, dict):
        value = payload.get("items")
        if isinstance(value, list):
            return [
                item
                for item in value
                if isinstance(item, dict)
            ]

    return []


def _flatten_text(value: Any) -> list[str]:
    result: list[str] = []

    if isinstance(value, str):
        cleaned = value.strip()
        if cleaned:
            result.append(cleaned)

    elif isinstance(value, dict):
        preferred_keys = (
            "text",
            "content",
            "markdown",
            "value",
            "recognized_text",
        )

        for key in preferred_keys:
            if key in value:
                result.extend(_flatten_text(value[key]))

        for key, child in value.items():
            if key not in preferred_keys:
                result.extend(_flatten_text(child))

    elif isinstance(value, list):
        for child in value:
            result.extend(_flatten_text(child))

    return result


def _decode_text(data: bytes) -> str:
    for encoding in (
        "utf-8-sig",
        "utf-8",
        "cp1252",
        "latin-1",
    ):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue

    return data.decode("utf-8", errors="replace")


def _extract_csv(data: bytes) -> str:
    text = _decode_text(data)
    rows = list(csv.reader(io.StringIO(text)))

    output = []
    for row in rows[:1000]:
        output.append(" | ".join(str(value) for value in row))

    return "\n".join(output)


def _extract_xlsx(data: bytes) -> str:
    workbook = load_workbook(
        io.BytesIO(data),
        read_only=True,
        data_only=True,
    )

    output: list[str] = []

    try:
        for sheet in workbook.worksheets[:20]:
            output.append(f"[SHEET] {sheet.title}")

            for row_index, row in enumerate(
                sheet.iter_rows(values_only=True),
                start=1,
            ):
                if row_index > 1000:
                    break

                values = [
                    "" if value is None else str(value)
                    for value in row
                ]

                if any(value.strip() for value in values):
                    output.append(" | ".join(values))

                if sum(len(line) for line in output) > MAX_EXTRACTED_CHARS:
                    break

            if sum(len(line) for line in output) > MAX_EXTRACTED_CHARS:
                break

    finally:
        workbook.close()

    return "\n".join(output)


def _extract_docx(data: bytes) -> str:
    document = Document(io.BytesIO(data))
    output: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            output.append(text)

    for table in document.tables:
        for row in table.rows:
            values = [
                cell.text.strip()
                for cell in row.cells
            ]
            if any(values):
                output.append(" | ".join(values))

    return "\n".join(output)


def _extract_pdf_text(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    output: list[str] = []

    for page in reader.pages[:100]:
        text = str(page.extract_text() or "").strip()
        if text:
            output.append(text)

        if sum(len(line) for line in output) > MAX_EXTRACTED_CHARS:
            break

    return "\n\n".join(output)


def _document_assistant_json(
    *,
    filename: str,
    content_type: str,
    data: bytes,
) -> tuple[str, dict[str, Any]]:
    response = httpx.post(
        DOCUMENT_ASSISTANT_URL,
        files={
            "file": (
                filename,
                data,
                content_type,
            )
        },
        data={
            "output_format": "json",
            "preprocess": "light",
            "dpi": "250",
            "min_confidence": "0.50",
            "max_pages": "100",
        },
        timeout=httpx.Timeout(
            1200.0,
            connect=10.0,
        ),
        follow_redirects=True,
    )

    if response.status_code >= 400:
        try:
            body = response.json()
            detail = body.get("detail", body)
            if isinstance(detail, dict):
                detail = detail.get(
                    "message",
                    str(detail),
                )
        except Exception:
            detail = response.text

        raise RuntimeError(
            f"OCR Document Assistant gagal: {detail}"
        )

    content = response.content
    payload: Any = None

    if zipfile.is_zipfile(io.BytesIO(content)):
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            json_names = [
                name
                for name in archive.namelist()
                if name.lower().endswith(".json")
            ]

            if not json_names:
                raise RuntimeError(
                    "Hasil OCR tidak memiliki JSON."
                )

            preferred = sorted(
                json_names,
                key=lambda name: (
                    "manifest" in name.lower(),
                    len(name),
                ),
            )[0]

            payload = json.loads(
                archive.read(preferred).decode(
                    "utf-8",
                    errors="replace",
                )
            )
    else:
        try:
            payload = json.loads(
                content.decode(
                    "utf-8",
                    errors="replace",
                )
            )
        except Exception as exc:
            raise RuntimeError(
                "Hasil OCR bukan JSON yang dapat dibaca."
            ) from exc

    text_parts = _flatten_text(payload)
    text = "\n".join(text_parts)

    if not text.strip():
        raise RuntimeError(
            "Tidak ada teks yang berhasil dikenali."
        )

    metadata = {
        "engine": "document-assistant-convert",
        "source": "ocr",
    }

    if isinstance(payload, dict):
        for key in (
            "page_count",
            "text_blocks",
            "low_confidence_blocks",
            "elapsed_seconds",
        ):
            if key in payload:
                metadata[key] = payload[key]

    return text, metadata


def _extract_upload(
    *,
    filename: str,
    content_type: str,
    data: bytes,
) -> tuple[str, dict[str, Any]]:
    suffix = Path(filename).suffix.lower()

    if suffix in {
        ".txt",
        ".md",
        ".log",
    }:
        return _decode_text(data), {
            "engine": "direct-text",
        }

    if suffix == ".csv":
        return _extract_csv(data), {
            "engine": "csv",
        }

    if suffix == ".json":
        payload = json.loads(
            _decode_text(data)
        )
        return json.dumps(
            payload,
            indent=2,
            ensure_ascii=False,
        ), {
            "engine": "json",
        }

    if suffix in {
        ".xlsx",
        ".xlsm",
    }:
        return _extract_xlsx(data), {
            "engine": "openpyxl",
        }

    if suffix == ".docx":
        return _extract_docx(data), {
            "engine": "python-docx",
        }

    if suffix == ".pdf":
        text = _extract_pdf_text(data)

        if len(text.strip()) >= 20:
            return text, {
                "engine": "pypdf",
            }

        return _document_assistant_json(
            filename=filename,
            content_type=(
                content_type
                or "application/pdf"
            ),
            data=data,
        )

    if suffix in {
        ".png",
        ".jpg",
        ".jpeg",
        ".webp",
        ".tif",
        ".tiff",
        ".bmp",
    }:
        return _document_assistant_json(
            filename=filename,
            content_type=(
                content_type
                or mimetypes.guess_type(
                    filename
                )[0]
                or "application/octet-stream"
            ),
            data=data,
        )

    raise ValueError(
        "Format belum didukung. Gunakan PDF, gambar, "
        "Excel, CSV, JSON, TXT, Markdown, atau DOCX."
    )


def register_smartbiz_ai_studio(
    *,
    app: Any,
    templates: Any,
    settings: Any,
    current_user: Callable[
        ...,
        dict[str, Any] | None,
    ],
    core_api_request: Callable[
        ...,
        Any,
    ],
    CoreAPIError: type[Exception],
    base_context: Callable[
        [Request],
        dict[str, Any],
    ],
) -> None:
    def login_redirect() -> RedirectResponse:
        return RedirectResponse(
            "/login?product=business",
            status_code=303,
        )

    def require_platform_csrf(
        request: Request,
    ) -> None:
        expected = str(
            request.session.get(
                "core_csrf_token",
                "",
            )
            or ""
        )

        received = str(
            request.headers.get(
                "X-CSRF-Token",
                "",
            )
            or ""
        )

        if (
            not expected
            or not received
            or expected != received
        ):
            raise PermissionError(
                "Token CSRF tidak valid."
            )

    def core_request(
        request: Request,
        method: str,
        path: str,
        *,
        params: dict[str, Any]
        | None = None,
        payload: dict[str, Any]
        | None = None,
        csrf: bool = False,
    ) -> Any:
        kwargs: dict[str, Any] = {}

        if params is not None:
            kwargs["params"] = params

        if payload is not None:
            kwargs["json_data"] = payload

        if csrf:
            kwargs["csrf"] = True

        response = core_api_request(
            request,
            settings,
            method,
            path,
            **kwargs,
        )

        return response.json()

    def core_get(
        request: Request,
        path: str,
        *,
        params: dict[str, Any]
        | None = None,
    ) -> Any:
        return core_request(
            request,
            "GET",
            path,
            params=params,
        )

    def core_post(
        request: Request,
        path: str,
        payload: dict[str, Any],
    ) -> Any:
        return core_request(
            request,
            "POST",
            path,
            payload=payload,
            csrf=True,
        )

    def load_businesses(
        request: Request,
    ) -> list[dict[str, Any]]:
        return _items(
            core_get(
                request,
                "/v1/smartbiz/me/businesses",
            )
        )

    def load_state(
        request: Request,
        business_id: str,
    ) -> dict[str, Any]:
        prefix = (
            "/v1/smartbiz/me/businesses/"
            f"{business_id}"
        )

        return {
            "business_id": business_id,
            "dashboard": core_get(
                request,
                f"{prefix}/dashboard",
            ),
            "products": _items(
                core_get(
                    request,
                    f"{prefix}/products",
                    params={
                        "limit": 200,
                        "offset": 0,
                    },
                )
            ),
            "customers": _items(
                core_get(
                    request,
                    f"{prefix}/customers",
                    params={
                        "limit": 200,
                        "offset": 0,
                    },
                )
            ),
            "locations": _items(
                core_get(
                    request,
                    f"{prefix}/locations",
                )
            ),
            "stock": _items(
                core_get(
                    request,
                    f"{prefix}/stock",
                    params={
                        "limit": 300,
                        "offset": 0,
                    },
                )
            ),
            "orders": _items(
                core_get(
                    request,
                    f"{prefix}/orders",
                    params={
                        "limit": 100,
                        "offset": 0,
                    },
                )
            ),
            "invoices": _items(
                core_get(
                    request,
                    f"{prefix}/invoices",
                    params={
                        "limit": 100,
                        "offset": 0,
                    },
                )
            ),
            "payments": _items(
                core_get(
                    request,
                    f"{prefix}/payments",
                    params={
                        "limit": 100,
                        "offset": 0,
                    },
                )
            ),
            "action_requests": _items(
                core_get(
                    request,
                    f"{prefix}/action-requests",
                    params={
                        "limit": 100,
                    },
                )
            ),
        }

    def ai_context(
        state: dict[str, Any],
    ) -> str:
        compact = {
            "business_id": state.get(
                "business_id"
            ),
            "dashboard": state.get(
                "dashboard"
            ),
            "products": (
                state.get("products") or []
            )[:180],
            "customers": (
                state.get("customers") or []
            )[:180],
            "locations": (
                state.get("locations") or []
            )[:20],
            "stock": (
                state.get("stock") or []
            )[:250],
            "orders": (
                state.get("orders") or []
            )[:80],
            "invoices": (
                state.get("invoices") or []
            )[:80],
            "payments": (
                state.get("payments") or []
            )[:80],
        }

        return json.dumps(
            compact,
            ensure_ascii=False,
            default=str,
            separators=(",", ":"),
        )[:30000]

    def call_ai(
        *,
        business_context: str,
        command: str,
        source: str,
        selected_context: dict[
            str,
            Any,
        ],
        attachment_name: str,
        attachment_type: str,
        attachment_text: str,
    ) -> dict[str, Any]:
        enabled = bool(
            getattr(
                settings,
                "business_ai_enabled",
                False,
            )
        )

        token = str(
            getattr(
                settings,
                "business_ai_token",
                "",
            )
            or ""
        ).strip()

        base_url = str(
            getattr(
                settings,
                "business_ai_url",
                (
                    "http://"
                    "botconnector-personal-ai:"
                    "8000"
                ),
            )
        ).rstrip("/")

        timeout = float(
            getattr(
                settings,
                "business_ai_timeout_seconds",
                100.0,
            )
        )

        if not enabled:
            raise RuntimeError(
                "Business AI belum aktif."
            )

        if not token:
            raise RuntimeError(
                "Token internal Business AI belum tersedia."
            )

        response = httpx.post(
            (
                f"{base_url}/v1/business/"
                "smartbiz/studio/plan"
            ),
            headers={
                "Authorization": (
                    f"Bearer {token}"
                ),
            },
            json={
                "business_context": (
                    business_context
                ),
                "command": command,
                "source": source,
                "selected_context": (
                    selected_context
                ),
                "attachment_name": (
                    attachment_name
                ),
                "attachment_type": (
                    attachment_type
                ),
                "attachment_text": (
                    attachment_text
                ),
            },
            timeout=httpx.Timeout(
                timeout,
                connect=5.0,
            ),
        )

        if response.status_code >= 400:
            try:
                detail = response.json().get(
                    "detail",
                    response.text,
                )
            except Exception:
                detail = response.text

            raise RuntimeError(
                f"SmartBiz AI gagal: {detail}"
            )

        payload = response.json()
        plan = payload.get("plan")

        if not isinstance(plan, dict):
            raise RuntimeError(
                "Struktur hasil SmartBiz AI tidak lengkap."
            )

        return plan

    @app.get(
        "/smartbiz",
        include_in_schema=False,
        response_class=HTMLResponse,
    )
    def smartbiz_studio_home(
        request: Request,
    ) -> Any:
        user = current_user(request)

        if not user:
            return login_redirect()

        try:
            businesses = load_businesses(
                request
            )
        except CoreAPIError as exc:
            status_code = int(
                getattr(
                    exc,
                    "status_code",
                    502,
                )
            )

            if status_code == 403:
                return RedirectResponse(
                    "/onboarding/business",
                    status_code=303,
                )

            return JSONResponse(
                {
                    "detail": str(
                        getattr(
                            exc,
                            "detail",
                            str(exc),
                        )
                    )
                },
                status_code=(
                    status_code
                    if status_code < 500
                    else 502
                ),
            )

        business_id = str(
            request.query_params.get(
                "business_id",
                "",
            )
        ).strip()

        if (
            not business_id
            and businesses
        ):
            business_id = str(
                businesses[0].get(
                    "id",
                    "",
                )
            )

        context = base_context(request)

        context.update(
            {
                "smartbiz_businesses": businesses,
                "selected_business_id": business_id,
                "csrf_token": str(
                    request.session.get(
                        "core_csrf_token",
                        "",
                    )
                    or ""
                ),
                "smartbiz_studio_version": "0.7.5",
                "user": user,
            }
        )

        return templates.TemplateResponse(
            request,
            "smartbiz_ai_studio.html",
            context,
        )

    @app.get(
        "/smartbiz/studio/state",
        include_in_schema=False,
    )
    def smartbiz_studio_state(
        request: Request,
        business_id: str,
    ) -> JSONResponse:
        if not current_user(request):
            return JSONResponse(
                {
                    "detail": "Sesi login diperlukan."
                },
                status_code=401,
            )

        try:
            businesses = load_businesses(
                request
            )

            business = next(
                (
                    item
                    for item in businesses
                    if str(item.get("id"))
                    == business_id
                ),
                None,
            )

            if business is None:
                return JSONResponse(
                    {
                        "detail": "Usaha tidak ditemukan."
                    },
                    status_code=404,
                )

            state = load_state(
                request,
                business_id,
            )

            state["business"] = business

            return JSONResponse(state)

        except CoreAPIError as exc:
            return JSONResponse(
                {
                    "detail": str(
                        getattr(
                            exc,
                            "detail",
                            str(exc),
                        )
                    )
                },
                status_code=int(
                    getattr(
                        exc,
                        "status_code",
                        400,
                    )
                ),
            )

        except Exception as exc:
            return JSONResponse(
                {
                    "detail": str(exc),
                },
                status_code=500,
            )

    @app.post(
        "/smartbiz/studio/upload",
        include_in_schema=False,
    )
    async def smartbiz_studio_upload(
        request: Request,
        file: UploadFile = File(...),
    ) -> JSONResponse:
        if not current_user(request):
            return JSONResponse(
                {
                    "detail": "Sesi login diperlukan."
                },
                status_code=401,
            )

        try:
            require_platform_csrf(request)

            filename = str(
                file.filename or "lampiran"
            )[:255]

            content_type = str(
                file.content_type
                or mimetypes.guess_type(
                    filename
                )[0]
                or "application/octet-stream"
            )

            data = await file.read(
                MAX_UPLOAD_BYTES + 1
            )

            if len(data) > MAX_UPLOAD_BYTES:
                raise ValueError(
                    "Ukuran lampiran maksimal 25 MB."
                )

            text, metadata = _extract_upload(
                filename=filename,
                content_type=content_type,
                data=data,
            )

            text = text.strip()

            if not text:
                raise ValueError(
                    "Lampiran tidak memiliki teks yang dapat dibaca."
                )

            return JSONResponse(
                {
                    "ok": True,
                    "filename": filename,
                    "content_type": content_type,
                    "size_bytes": len(data),
                    "text": text[
                        :MAX_EXTRACTED_CHARS
                    ],
                    "truncated": (
                        len(text)
                        > MAX_EXTRACTED_CHARS
                    ),
                    "metadata": metadata,
                }
            )

        except PermissionError as exc:
            return JSONResponse(
                {
                    "detail": str(exc),
                },
                status_code=403,
            )

        except ValueError as exc:
            return JSONResponse(
                {
                    "detail": str(exc),
                },
                status_code=422,
            )

        except Exception as exc:
            return JSONResponse(
                {
                    "detail": str(exc),
                },
                status_code=500,
            )

    @app.post(
        "/smartbiz/studio/command",
        include_in_schema=False,
    )
    async def smartbiz_studio_command(
        request: Request,
    ) -> JSONResponse:
        user = current_user(request)

        if not user:
            return JSONResponse(
                {
                    "detail": "Sesi login diperlukan."
                },
                status_code=401,
            )

        try:
            require_platform_csrf(request)
            body = await request.json()
        except PermissionError as exc:
            return JSONResponse(
                {
                    "detail": str(exc),
                },
                status_code=403,
            )
        except Exception:
            body = {}

        business_id = str(
            body.get("business_id", "")
        ).strip()

        command = str(
            body.get("command", "")
        ).strip()

        source = str(
            body.get("source", "web")
        ).strip().lower()

        selected_context = body.get(
            "selected_context"
        )

        attachment_name = str(
            body.get(
                "attachment_name",
                "",
            )
        )[:255]

        attachment_type = str(
            body.get(
                "attachment_type",
                "",
            )
        )[:100]

        attachment_text = str(
            body.get(
                "attachment_text",
                "",
            )
        )[:MAX_EXTRACTED_CHARS]

        if source not in {
            "web",
            "voice",
            "barcode",
            "document",
        }:
            source = "web"

        if not isinstance(
            selected_context,
            dict,
        ):
            selected_context = {}

        if not business_id:
            return JSONResponse(
                {
                    "detail": "Pilih usaha terlebih dahulu."
                },
                status_code=422,
            )

        if len(command) < 2:
            return JSONResponse(
                {
                    "detail": "Tuliskan kebutuhan usaha."
                },
                status_code=422,
            )

        try:
            state = load_state(
                request,
                business_id,
            )

            plan = call_ai(
                business_context=ai_context(
                    state
                ),
                command=command,
                source=source,
                selected_context=(
                    selected_context
                ),
                attachment_name=(
                    attachment_name
                ),
                attachment_type=(
                    attachment_type
                ),
                attachment_text=(
                    attachment_text
                ),
            )

            actions = plan.get("actions")
            if not isinstance(actions, list):
                actions = []

            results: list[dict[str, Any]] = []

            for action in actions:
                if not isinstance(action, dict):
                    continue

                action_type = str(
                    action.get(
                        "action_type",
                        "",
                    )
                ).strip()

                payload = action.get("payload")
                if not isinstance(payload, dict):
                    payload = {}

                confidence = float(
                    action.get(
                        "confidence",
                        0,
                    )
                    or 0
                )

                requires_approval = bool(
                    action.get(
                        "requires_approval",
                        True,
                    )
                )

                if action_type == "check_stock":
                    results.append(
                        {
                            **action,
                            "status": (
                                "completed_read_only"
                            ),
                            "result": state.get(
                                "stock"
                            ),
                        }
                    )
                    continue

                auto_path = None

                if (
                    not requires_approval
                    and confidence >= 0.95
                ):
                    if action_type == "create_customer":
                        auto_path = (
                            "/v1/smartbiz/me/"
                            "businesses/"
                            f"{business_id}/customers"
                        )

                    elif action_type == "create_product":
                        auto_path = (
                            "/v1/smartbiz/me/"
                            "businesses/"
                            f"{business_id}/products"
                        )

                    elif action_type == "create_order_draft":
                        auto_path = (
                            "/v1/smartbiz/me/"
                            "businesses/"
                            f"{business_id}/orders/draft"
                        )

                        payload.setdefault(
                            "idempotency_key",
                            (
                                "studio-order-"
                                + uuid4().hex
                            ),
                        )

                if auto_path:
                    created = core_post(
                        request,
                        auto_path,
                        payload,
                    )

                    results.append(
                        {
                            **action,
                            "status": (
                                "completed_automatically"
                            ),
                            "result": created,
                        }
                    )
                    continue

                proposal = core_post(
                    request,
                    (
                        "/v1/smartbiz/me/"
                        "businesses/"
                        f"{business_id}/ai-actions"
                    ),
                    {
                        "action_type": action_type,
                        "title": str(
                            action.get(
                                "title",
                                action_type,
                            )
                        ),
                        "explanation": str(
                            action.get(
                                "explanation",
                                (
                                    "Tindakan memerlukan "
                                    "persetujuan."
                                ),
                            )
                        ),
                        "risk_level": str(
                            action.get(
                                "risk_level",
                                "medium",
                            )
                        ),
                        "confidence": confidence,
                        "extracted_fields": (
                            action.get(
                                "extracted_fields",
                                {},
                            )
                            if isinstance(
                                action.get(
                                    "extracted_fields"
                                ),
                                dict,
                            )
                            else {}
                        ),
                        "payload": payload,
                        "idempotency_key": (
                            "studio-action-"
                            + uuid4().hex
                        ),
                    },
                )

                results.append(
                    {
                        **action,
                        "status": (
                            "approval_required"
                        ),
                        "request": proposal,
                    }
                )

            updated_plan = dict(plan)
            updated_plan["actions"] = results

            return JSONResponse(
                {
                    "ok": True,
                    "operational": True,
                    "autofill": True,
                    "shadow_mode": False,
                    "plan": updated_plan,
                }
            )

        except CoreAPIError as exc:
            return JSONResponse(
                {
                    "detail": str(
                        getattr(
                            exc,
                            "detail",
                            str(exc),
                        )
                    )
                },
                status_code=int(
                    getattr(
                        exc,
                        "status_code",
                        400,
                    )
                ),
            )

        except Exception as exc:
            return JSONResponse(
                {
                    "detail": str(exc),
                },
                status_code=500,
            )

    @app.post(
        (
            "/smartbiz/studio/actions/"
            "{request_id}/review"
        ),
        include_in_schema=False,
    )
    async def smartbiz_studio_review(
        request: Request,
        request_id: str,
    ) -> JSONResponse:
        if not current_user(request):
            return JSONResponse(
                {
                    "detail": "Sesi login diperlukan."
                },
                status_code=401,
            )

        try:
            require_platform_csrf(request)
            body = await request.json()
        except PermissionError as exc:
            return JSONResponse(
                {
                    "detail": str(exc),
                },
                status_code=403,
            )
        except Exception:
            body = {}

        business_id = str(
            body.get("business_id", "")
        ).strip()

        decision = str(
            body.get("decision", "")
        ).strip().lower()

        notes = str(
            body.get("notes", "")
        ).strip() or None

        if decision not in {
            "approve",
            "reject",
        }:
            return JSONResponse(
                {
                    "detail": (
                        "Keputusan harus approve "
                        "atau reject."
                    )
                },
                status_code=422,
            )

        try:
            review = core_post(
                request,
                (
                    "/v1/smartbiz/me/"
                    "businesses/"
                    f"{business_id}/"
                    "action-requests/"
                    f"{request_id}/review"
                ),
                {
                    "decision": decision,
                    "notes": notes,
                },
            )

            execution = None

            if decision == "approve":
                execution = core_post(
                    request,
                    (
                        "/v1/smartbiz/me/"
                        "businesses/"
                        f"{business_id}/"
                        "ai-actions/"
                        f"{request_id}/execute"
                    ),
                    {},
                )

            return JSONResponse(
                {
                    "ok": True,
                    "review": review,
                    "execution": execution,
                }
            )

        except CoreAPIError as exc:
            return JSONResponse(
                {
                    "detail": str(
                        getattr(
                            exc,
                            "detail",
                            str(exc),
                        )
                    )
                },
                status_code=int(
                    getattr(
                        exc,
                        "status_code",
                        400,
                    )
                ),
            )

        except Exception as exc:
            return JSONResponse(
                {
                    "detail": str(exc),
                },
                status_code=500,
            )
