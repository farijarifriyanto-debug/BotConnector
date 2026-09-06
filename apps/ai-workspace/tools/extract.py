"""Extract plain text and safe metadata from Study AI uploads."""
from __future__ import annotations

import io
from typing import Any

MAX_CHARS = 60_000


def _cap(text: str) -> tuple[str, bool]:
    text = text.strip()
    return text[:MAX_CHARS], len(text) > MAX_CHARS


def extract_document(name: str, data: bytes) -> dict[str, Any]:
    ext = (name.rsplit(".", 1)[-1] if "." in name else "").lower()

    if ext in {"txt", "md", "csv"}:
        text, truncated = _cap(data.decode("utf-8", errors="ignore"))
        if not text:
            raise ValueError("Dokumen kosong / tanpa teks.")
        return {
            "text": text,
            "kind": ext,
            "pages": None,
            "truncated": truncated,
        }

    if ext == "pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        parts: list[str] = []
        total = 0
        truncated = False
        pages_read = 0
        extracted_body_chars = 0

        for page_number, page in enumerate(reader.pages, start=1):
            body = (page.extract_text() or "").strip()
            extracted_body_chars += len(body)
            marker = f"\n\n--- Halaman {page_number} ---\n"
            page_text = marker + body
            remaining = MAX_CHARS - total
            if remaining <= 0:
                truncated = True
                break
            if len(page_text) > remaining:
                parts.append(page_text[:remaining])
                total += remaining
                truncated = True
                pages_read = page_number
                break
            parts.append(page_text)
            total += len(page_text)
            pages_read = page_number

        text = "".join(parts).strip()
        if not text or extracted_body_chars == 0:
            raise ValueError(
                "PDF ini sepertinya hasil scan (tanpa teks). "
                "Pakai Document Assistant untuk OCR dulu."
            )
        if pages_read < len(reader.pages):
            truncated = True
        return {
            "text": text,
            "kind": "pdf",
            "pages": len(reader.pages),
            "truncated": truncated,
        }

    if ext == "docx":
        import docx

        document = docx.Document(io.BytesIO(data))
        blocks: list[str] = []
        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                blocks.append(paragraph.text)
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    blocks.append(" | ".join(cells))
        text, truncated = _cap("\n".join(blocks))
        if not text:
            raise ValueError("Dokumen kosong / tanpa teks.")
        return {
            "text": text,
            "kind": "docx",
            "pages": None,
            "truncated": truncated,
        }

    if ext == "doc":
        raise ValueError("Format .doc lama belum didukung — simpan sebagai .docx atau PDF.")

    raise ValueError(
        f"Format .{ext or '?'} tidak didukung "
        "(pakai PDF, DOCX, TXT, Markdown, atau CSV)."
    )


def extract_text(name: str, data: bytes) -> str:
    """Backward-compatible helper retained for older callers."""
    return str(extract_document(name, data)["text"])
